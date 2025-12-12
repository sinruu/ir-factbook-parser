"""
메인 파서 모듈

4대 금융지주 IR FactBook Excel 파일을 파싱하여 정규화된 데이터로 변환
"""

import re
import yaml
import json
import csv
from pathlib import Path
from typing import Dict, List, Optional, Any, Iterator
from dataclasses import dataclass, field
import logging

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
import pandas as pd

from .period_parser import PeriodParser, ParsedPeriod, parse_period, is_skip_column
from .normalizer import Normalizer, NormalizedRecord, get_normalizer
from .validators import DataValidator, ValidationReport, validate_records


# 로깅 설정
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


@dataclass
class SheetConfig:
    """시트 파싱 설정"""
    entity: str
    category: str
    header_row: int
    data_start_row: int
    metric_columns: List[int]
    data_start_col: int
    default_unit: str = "KRW_billion"
    metrics_override: List[Dict[str, str]] = field(default_factory=list)
    has_annual_columns: bool = False
    cumulative_data: bool = False


@dataclass
class CompanyConfig:
    """회사별 파싱 설정"""
    name: str
    name_en: str
    code: str
    file_patterns: List[str]
    sheets: Dict[str, SheetConfig]
    skip_sheets: List[str]
    special_rules: Dict[str, Any]


class FactBookParser:
    """IR FactBook 파서"""

    def __init__(self, config_dir: Optional[Path] = None):
        """
        Args:
            config_dir: 설정 파일 디렉토리 경로
        """
        if config_dir is None:
            config_dir = Path(__file__).parent.parent / "config"
        self.config_dir = Path(config_dir)

        self.period_parser = PeriodParser()
        self.normalizer = get_normalizer(config_dir)

        self.company_configs: Dict[str, CompanyConfig] = {}
        self._load_configs()

    def _load_configs(self):
        """회사별 설정 파일 로드"""
        for config_file in self.config_dir.glob("*.yaml"):
            if config_file.name == "schema.yaml":
                continue

            try:
                with open(config_file, "r", encoding="utf-8") as f:
                    config = yaml.safe_load(f)

                if "company" not in config:
                    continue

                company_code = config["company"]["code"]
                self.company_configs[company_code] = self._parse_company_config(config)
                logger.info(f"Loaded config for {company_code}")

            except Exception as e:
                logger.warning(f"Failed to load config {config_file}: {e}")

    def _parse_company_config(self, config: Dict) -> CompanyConfig:
        """회사 설정 파싱"""
        sheets = {}
        for sheet_name, sheet_config in config.get("sheets", {}).items():
            structure = sheet_config.get("structure", {})
            sheets[sheet_name] = SheetConfig(
                entity=sheet_config.get("entity", "지주"),
                category=sheet_config.get("category", "Other"),
                header_row=structure.get("header_row", 3),
                data_start_row=structure.get("data_start_row", 5),
                metric_columns=structure.get("metric_columns", [3]),
                data_start_col=structure.get("data_start_col", 7),
                default_unit=sheet_config.get("default_unit", "KRW_billion"),
                metrics_override=sheet_config.get("metrics_override", []),
                has_annual_columns=structure.get("has_annual_columns", False),
                cumulative_data=sheet_config.get("special_handling", {}).get("cumulative_data", False)
            )

        return CompanyConfig(
            name=config["company"]["name"],
            name_en=config["company"]["name_en"],
            code=config["company"]["code"],
            file_patterns=config.get("file_patterns", []),
            sheets=sheets,
            skip_sheets=config.get("skip_sheets", []),
            special_rules=config.get("special_rules", {})
        )

    def detect_company(self, file_path: Path) -> Optional[str]:
        """파일명으로 회사 감지"""
        filename = file_path.name.lower()

        # 파일 패턴 매칭
        for code, config in self.company_configs.items():
            for pattern in config.file_patterns:
                # 간단한 와일드카드 패턴 -> 정규식
                regex_pattern = pattern.replace("*", ".*").lower()
                if re.search(regex_pattern, filename):
                    return code

        # 키워드 기반 감지
        if "woori" in filename or "우리" in filename:
            return "woori"
        elif "kb" in filename:
            return "kb"
        elif "sfg" in filename or "shinhan" in filename or "신한" in filename:
            return "shinhan"
        elif "hfg" in filename or "hana" in filename or "하나" in filename:
            return "hana"

        return None

    def parse_file(
        self,
        file_path: Path,
        company_code: Optional[str] = None
    ) -> List[NormalizedRecord]:
        """
        Excel 파일 파싱

        Args:
            file_path: Excel 파일 경로
            company_code: 회사 코드 (None이면 자동 감지)

        Returns:
            NormalizedRecord 목록
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # 회사 감지
        if company_code is None:
            company_code = self.detect_company(file_path)
            if company_code is None:
                raise ValueError(f"Cannot detect company from filename: {file_path.name}")

        logger.info(f"Parsing {file_path.name} as {company_code}")

        # 회사 설정 가져오기
        config = self.company_configs.get(company_code)
        if config is None:
            logger.warning(f"No config for {company_code}, using default parsing")
            return self._parse_generic(file_path, company_code)

        # Excel 파일 열기
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        records = []

        for sheet_name in workbook.sheetnames:
            # 스킵할 시트인지 확인
            if sheet_name in config.skip_sheets:
                logger.debug(f"Skipping sheet: {sheet_name}")
                continue

            # 시트 설정 찾기
            sheet_config = self._find_sheet_config(sheet_name, config)
            if sheet_config is None:
                logger.debug(f"No config for sheet: {sheet_name}")
                continue

            try:
                sheet = workbook[sheet_name]
                sheet_records = self._parse_sheet(
                    sheet, sheet_name, sheet_config, company_code
                )
                records.extend(sheet_records)
                logger.info(f"Parsed {len(sheet_records)} records from {sheet_name}")
            except Exception as e:
                logger.warning(f"Failed to parse sheet {sheet_name}: {e}")

        workbook.close()
        return records

    def _find_sheet_config(
        self,
        sheet_name: str,
        config: CompanyConfig
    ) -> Optional[SheetConfig]:
        """시트 설정 찾기 (정확히 일치하거나 부분 일치)"""
        # 정확히 일치
        if sheet_name in config.sheets:
            return config.sheets[sheet_name]

        # 부분 일치 (시트명에 공백이나 특수문자가 있을 수 있음)
        for config_name, sheet_config in config.sheets.items():
            if config_name.strip() == sheet_name.strip():
                return sheet_config
            # 부분 문자열 매칭
            if config_name in sheet_name or sheet_name in config_name:
                return sheet_config

        return None

    def _parse_sheet(
        self,
        sheet: Worksheet,
        sheet_name: str,
        config: SheetConfig,
        company_code: str
    ) -> List[NormalizedRecord]:
        """시트 파싱"""
        records = []

        # 헤더 행에서 기간 컬럼 파싱
        period_columns = self._parse_header_periods(
            sheet, config.header_row, config.data_start_col
        )

        if not period_columns:
            logger.warning(f"No period columns found in {sheet_name}")
            return records

        # 단위 오버라이드 맵 구축
        unit_override = {}
        for override in config.metrics_override:
            if "original" in override and "unit" in override:
                unit_override[override["original"].lower()] = override["unit"]

        # 데이터 행 순회
        for row_idx in range(config.data_start_row + 1, sheet.max_row + 1):
            # 지표명 추출 (계층적 열에서)
            metric_name = self._extract_metric_name(
                sheet, row_idx, config.metric_columns
            )

            if not metric_name:
                continue

            # 단위 결정
            unit = config.default_unit
            metric_lower = metric_name.lower()
            for override_key, override_unit in unit_override.items():
                if override_key in metric_lower:
                    unit = override_unit
                    break

            # 각 기간 컬럼의 값 추출
            for col_idx, parsed_period in period_columns.items():
                cell = sheet.cell(row=row_idx, column=col_idx)
                value = cell.value

                if value is None:
                    continue

                # 정규화된 레코드 생성
                record = self.normalizer.normalize_record(
                    company_code=company_code,
                    entity=config.entity,
                    category=config.category,
                    metric_original=metric_name,
                    period=parsed_period.standard_format,
                    year=parsed_period.year,
                    quarter=parsed_period.quarter,
                    value=value,
                    unit=unit,
                    is_cumulative=config.cumulative_data or parsed_period.is_cumulative,
                    is_estimate=parsed_period.is_estimate
                )

                if record.value is not None:
                    records.append(record)

        return records

    def _parse_header_periods(
        self,
        sheet: Worksheet,
        header_row: int,
        start_col: int
    ) -> Dict[int, ParsedPeriod]:
        """헤더 행에서 기간 컬럼 파싱"""
        period_columns = {}

        for col_idx in range(start_col + 1, sheet.max_column + 1):
            cell = sheet.cell(row=header_row + 1, column=col_idx)
            header_value = cell.value

            if header_value is None:
                continue

            header_str = str(header_value).strip()

            # 스킵할 컬럼인지 확인
            if is_skip_column(header_str):
                continue

            # 기간 파싱
            parsed = parse_period(header_str)
            if parsed:
                period_columns[col_idx] = parsed

        return period_columns

    def _extract_metric_name(
        self,
        sheet: Worksheet,
        row_idx: int,
        metric_columns: List[int]
    ) -> Optional[str]:
        """지표명 추출 (계층적 열에서 마지막 비어있지 않은 값)"""
        parts = []

        for col_idx in metric_columns:
            cell = sheet.cell(row=row_idx, column=col_idx + 1)  # 1-indexed
            value = cell.value

            if value is not None:
                value_str = str(value).strip()
                if value_str:
                    parts.append(value_str)

        if not parts:
            return None

        # 마지막 비어있지 않은 값 반환 (가장 구체적인 지표명)
        return parts[-1]

    def _parse_generic(
        self,
        file_path: Path,
        company_code: str
    ) -> List[NormalizedRecord]:
        """설정 없이 기본 파싱 (fallback)"""
        logger.warning(f"Using generic parsing for {company_code}")

        records = []
        workbook = openpyxl.load_workbook(file_path, data_only=True)

        for sheet_name in workbook.sheetnames:
            # Financial Highlights 시트 우선 시도
            if "Financial" not in sheet_name and "Highlight" not in sheet_name:
                continue

            sheet = workbook[sheet_name]
            # 기본 설정으로 파싱 시도
            try:
                config = SheetConfig(
                    entity="지주",
                    category="Financial Highlights",
                    header_row=3,
                    data_start_row=5,
                    metric_columns=[3],
                    data_start_col=7,
                    default_unit="KRW_billion"
                )
                sheet_records = self._parse_sheet(sheet, sheet_name, config, company_code)
                records.extend(sheet_records)
            except Exception as e:
                logger.warning(f"Generic parsing failed for {sheet_name}: {e}")

        workbook.close()
        return records


def parse_factbook(
    file_path: str,
    company_code: Optional[str] = None,
    config_dir: Optional[str] = None
) -> List[NormalizedRecord]:
    """
    FactBook 파일 파싱 (편의 함수)

    Args:
        file_path: Excel 파일 경로
        company_code: 회사 코드 (woori, kb, shinhan, hana)
        config_dir: 설정 디렉토리 경로

    Returns:
        NormalizedRecord 목록
    """
    config_path = Path(config_dir) if config_dir else None
    parser = FactBookParser(config_path)
    return parser.parse_file(Path(file_path), company_code)


def export_to_csv(
    records: List[NormalizedRecord],
    output_path: str,
    include_header: bool = True
):
    """
    레코드를 CSV로 내보내기

    Args:
        records: NormalizedRecord 목록
        output_path: 출력 파일 경로
        include_header: 헤더 포함 여부
    """
    if not records:
        logger.warning("No records to export")
        return

    fieldnames = [
        "holding_company", "entity", "category", "metric_std",
        "metric_original", "period", "year", "quarter", "value",
        "unit", "is_cumulative", "is_estimate"
    ]

    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        if include_header:
            writer.writeheader()
        for record in records:
            writer.writerow(record.to_dict())

    logger.info(f"Exported {len(records)} records to {output_path}")


def export_to_json(
    records: List[NormalizedRecord],
    output_path: str,
    indent: int = 2
):
    """
    레코드를 JSON으로 내보내기

    Args:
        records: NormalizedRecord 목록
        output_path: 출력 파일 경로
        indent: JSON 들여쓰기
    """
    data = {
        "metadata": {
            "total_records": len(records),
            "schema_version": "1.0"
        },
        "records": [r.to_dict() for r in records]
    }

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=indent)

    logger.info(f"Exported {len(records)} records to {output_path}")


# CLI 엔트리포인트
if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python parser.py <input_file> [output_file]")
        sys.exit(1)

    input_file = sys.argv[1]
    output_file = sys.argv[2] if len(sys.argv) > 2 else "output/normalized.csv"

    records = parse_factbook(input_file)
    print(f"Parsed {len(records)} records")

    # 검증
    report = validate_records(records)
    report.print_summary()

    # 내보내기
    export_to_csv(records, output_file)
